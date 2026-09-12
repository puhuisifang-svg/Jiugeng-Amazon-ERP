#!/usr/bin/env python3
"""Generate Sprint 0 Excel import templates and synthetic validation fixtures."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[3]
OUTPUT_DIR = ROOT / "deliverables" / "pm" / "excel_templates"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")


def sheet_setup(ws, headers: list[str], comments: dict[str, str] | None = None) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).coordinate}"
    ws.row_dimensions[1].height = 30
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(1, idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if comments and header in comments:
            cell.comment = Comment(comments[header], "九更产品经理")
        ws.column_dimensions[cell.column_letter].width = min(max(len(header) + 4, 14), 28)


def add_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_enum_validation(ws, header: str, values: list[str], end_row: int = 1001) -> None:
    headers = {cell.value: cell.column for cell in ws[1]}
    if header not in headers:
        return
    formula = '"' + ",".join(values) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "请选择字典允许的枚举值"
    dv.errorTitle = "枚举值无效"
    dv.prompt = "允许值：" + ", ".join(values)
    dv.promptTitle = header
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    col = ws.cell(1, headers[header]).column_letter
    dv.add(f"{col}2:{col}{end_row}")


def add_decimal_validation(ws, header: str, operator: str, formula1: str, formula2: str | None = None) -> None:
    headers = {cell.value: cell.column for cell in ws[1]}
    if header not in headers:
        return
    kwargs = {"type": "decimal", "operator": operator, "formula1": formula1, "allow_blank": True}
    if formula2 is not None:
        kwargs["formula2"] = formula2
    dv = DataValidation(**kwargs)
    dv.error = "数值超出允许范围"
    dv.errorTitle = "数值无效"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    col = ws.cell(1, headers[header]).column_letter
    dv.add(f"{col}2:{col}1001")


def add_instructions(wb: Workbook, template_name: str, import_order: str, keys: str, caveats: list[str]) -> None:
    ws = wb.active
    ws.title = "填写说明"
    rows = [
        ["模板", template_name],
        ["版本", "v1.0 / 2026-09-12"],
        ["用途", "Sprint 0 字段与校验评审；导入实现应以 FIELD_DICTIONARY_V1.md 为准。"],
        ["数据声明", "TEST-* 内容是人工构造的校验测试夹具，不是真实业务样本，不得用于业务验收。"],
        ["导入顺序", import_order],
        ["幂等/唯一键", keys],
        ["空值规则", "去除首尾空格后为空即空值；必填字段不得为空。"],
        ["错误返回", "失败必须返回工作表、行号、字段和原因；不得静默跳过。"],
    ]
    rows.extend([["注意", item] for item in caveats])
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws["B4"].fill = NOTE_FILL


def add_enum_sheet(wb: Workbook, groups: dict[str, list[str]]) -> None:
    ws = wb.create_sheet("枚举值")
    headers = list(groups)
    sheet_setup(ws, headers)
    for row_idx in range(max(len(values) for values in groups.values())):
        ws.append([groups[h][row_idx] if row_idx < len(groups[h]) else None for h in headers])


def add_machine_metadata(wb: Workbook, import_type: str, max_data_rows: int = 1000) -> None:
    """Add machine-readable identity used by upload compatibility checks."""
    ws = wb.create_sheet("_meta")
    rows = [
        ["schemaIdentifier", "jiugeng.amazon.erp.import"],
        ["importType", import_type],
        ["templateVersion", "1.0"],
        ["maxDataRows", max_data_rows],
    ]
    add_rows(ws, rows)
    ws.sheet_state = "veryHidden"


def save_and_check(wb: Workbook, filename: str, required_sheets: list[str]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    wb.save(path)
    checked = load_workbook(path, data_only=False)
    if "_meta" not in checked.sheetnames or checked["_meta"].sheet_state != "veryHidden":
        raise RuntimeError(f"{filename} missing veryHidden _meta sheet")
    metadata = {row[0].value: row[1].value for row in checked["_meta"].iter_rows(min_row=1, max_col=2)}
    if metadata.get("schemaIdentifier") != "jiugeng.amazon.erp.import" or metadata.get("templateVersion") != "1.0":
        raise RuntimeError(f"{filename} has invalid machine metadata: {metadata}")
    missing = [name for name in required_sheets if name not in checked.sheetnames]
    if missing:
        raise RuntimeError(f"{filename} missing sheets: {missing}")
    for sheet_name in required_sheets:
        if checked[sheet_name].max_row < 1 or checked[sheet_name].max_column < 1:
            raise RuntimeError(f"{filename}:{sheet_name} is empty")


def product_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "01 产品/SPU 导入模板",
        "本模板 → 02 内部 SKU → 04 辅料 → 05 SKU-辅料 / 03 Listing",
        "product_code（大小写不敏感全局唯一）",
        ["product_code 由业务手工提供，系统不得自动生成。", "internal_category 的受控词表待真实样本复核。"],
    )
    add_machine_metadata(wb, "PRODUCTS")
    headers = ["product_code", "name_zh", "name_en", "brand", "internal_category", "product_type", "description", "status"]
    comments = {
        "product_code": "必填；1-64 位；^[A-Za-z0-9][A-Za-z0-9_-]{0,63}$；全局唯一。",
        "status": "必填：DRAFT / ACTIVE / INACTIVE。",
    }
    ws = wb.create_sheet("导入模板")
    sheet_setup(ws, headers, comments)
    add_enum_validation(ws, "status", ["DRAFT", "ACTIVE", "INACTIVE"])
    good = wb.create_sheet("正确样例")
    sheet_setup(good, headers, comments)
    add_rows(good, [["TEST-P001", "测试产品一（非业务数据）", "Synthetic Product One", "TEST-BRAND", "TEST-CATEGORY", "TEST_TYPE", "仅用于校验", "ACTIVE"]])
    bad_headers = headers + ["_预期错误"]
    bad = wb.create_sheet("错误样例")
    sheet_setup(bad, bad_headers)
    add_rows(bad, [
        [None, "缺编码", None, None, "TEST-CATEGORY", None, None, "ACTIVE", "product_code 必填"],
        ["测试 产品", "非法编码", None, None, "TEST-CATEGORY", None, None, "ACTIVE", "编码含中文/空格"],
        ["TEST-P001", "重复编码 A", None, None, "TEST-CATEGORY", None, None, "ACTIVE", "与下一行大小写不敏感重复"],
        ["test-p001", "重复编码 B", None, None, "TEST-CATEGORY", None, None, "UNKNOWN", "编码重复；status 枚举无效"],
    ])
    add_enum_sheet(wb, {"status": ["DRAFT", "ACTIVE", "INACTIVE"]})
    save_and_check(wb, "01_products_v1.0.xlsx", ["填写说明", "导入模板", "正确样例", "错误样例", "枚举值"])


def sku_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "02 内部 SKU 与动态属性导入模板",
        "先导入 01 产品；本文件内先 variants，再 variant_attributes",
        "sku_code 全局唯一；variant_attributes: sku_code + attribute_code；同 Product 属性组合唯一",
        ["barcode 非空时全局唯一。", "动态属性使用规范化明细表，不限制属性轴数量。"],
    )
    add_machine_metadata(wb, "VARIANTS")
    variant_headers = ["sku_code", "product_code", "variant_name", "barcode", "material", "specification", "item_dimensions", "package_dimensions", "package_weight", "package_weight_unit", "status"]
    comments = {
        "sku_code": "必填；业务手工提供；编码规则同产品编码。",
        "product_code": "必填；引用已存在且未停用 Product。",
        "barcode": "可空；非空时全局唯一；仅 8/12/13/14 位数字。",
        "package_weight_unit": "有 package_weight 时必填。",
    }
    ws = wb.create_sheet("variants_导入模板")
    sheet_setup(ws, variant_headers, comments)
    add_enum_validation(ws, "package_weight_unit", ["G", "KG", "OZ", "LB"])
    add_enum_validation(ws, "status", ["ACTIVE", "INACTIVE"])
    add_decimal_validation(ws, "package_weight", "greaterThan", "0")
    attr_headers = ["sku_code", "attribute_code", "attribute_value", "display_order"]
    attrs = wb.create_sheet("attributes_导入模板")
    sheet_setup(attrs, attr_headers, {"sku_code": "引用本文件 variants 或库内 SKU。", "attribute_code": "同一 SKU 内唯一。", "display_order": ">=1，同一 SKU 内唯一。"})
    good = wb.create_sheet("正确样例_variants")
    sheet_setup(good, variant_headers)
    add_rows(good, [
        ["TEST-SKU001", "TEST-P001", "黑色大号测试规格", "123456789012", "TEST-MATERIAL", "TEST-SPEC", "10 x 20 x 5 cm", "12 x 22 x 7 cm", 0.5, "KG", "ACTIVE"],
        ["TEST-SKU002", "TEST-P001", "白色大号测试规格", "123456789013", "TEST-MATERIAL", "TEST-SPEC", "10 x 20 x 5 cm", "12 x 22 x 7 cm", 0.5, "KG", "ACTIVE"],
    ])
    good_attr = wb.create_sheet("正确样例_attributes")
    sheet_setup(good_attr, attr_headers)
    add_rows(good_attr, [
        ["TEST-SKU001", "color", "Black", 1], ["TEST-SKU001", "size", "Large", 2],
        ["TEST-SKU002", "color", "White", 1], ["TEST-SKU002", "size", "Large", 2],
    ])
    bad = wb.create_sheet("错误样例_variants")
    sheet_setup(bad, variant_headers + ["_预期错误"])
    add_rows(bad, [
        ["TEST-SKU-BAD1", "TEST-P-NOT-FOUND", "引用错误", None, None, None, None, None, None, None, "ACTIVE", "product_code 不存在"],
        ["TEST-SKU-BAD2", "TEST-P001", "条码错误", "ABC", None, None, None, None, -1, "STONE", "UNKNOWN", "barcode/重量/单位/status 无效"],
        ["TEST-SKU001", "TEST-P001", "重复 SKU", None, None, None, None, None, None, None, "ACTIVE", "sku_code 与正确样例重复"],
    ])
    bad_attr = wb.create_sheet("错误样例_attributes")
    sheet_setup(bad_attr, attr_headers + ["_预期错误"])
    add_rows(bad_attr, [
        ["TEST-SKU-NOT-FOUND", "color", "Black", 1, "sku_code 不存在"],
        ["TEST-SKU001", "color", "Black", 1, "与下一行 attribute_code/display_order 重复"],
        ["TEST-SKU001", "color", "White", 1, "同一 SKU attribute_code/display_order 重复"],
        ["TEST-SKU002", None, "Large", 0, "属性编码必填；display_order 必须 >=1"],
    ])
    add_enum_sheet(wb, {"package_weight_unit": ["G", "KG", "OZ", "LB"], "status": ["ACTIVE", "INACTIVE"]})
    save_and_check(wb, "02_variants_v1.0.xlsx", ["填写说明", "variants_导入模板", "attributes_导入模板", "正确样例_variants", "正确样例_attributes", "错误样例_variants", "错误样例_attributes", "枚举值"])


def listing_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "03 Amazon 美国站 5 店 Listing 映射导入模板",
        "先导入 01 产品与 02 SKU；父体行必须先于引用它的子体行",
        "store_code + marketplace_id + seller_sku；同 SKU 同店最多一个 ACTIVE Listing",
        ["首期 store_code 必须引用系统中预先配置的 5 个启用店铺。", "测试样例中的 TEST-STORE-* 仅为结构测试，不代表真实店铺。"],
    )
    add_machine_metadata(wb, "STORE_LISTINGS")
    headers = ["store_code", "marketplace_id", "seller_sku", "asin", "listing_role", "product_code", "sku_code", "parent_seller_sku", "variation_theme", "fnsku", "fulfillment_channel", "listing_status"]
    comments = {
        "marketplace_id": "美国站固定 ATVPDKIKX0DER。",
        "sku_code": "PARENT 必须为空；CHILD/STANDALONE 必填。",
        "parent_seller_sku": "CHILD 必填，且引用同店 PARENT。",
        "fnsku": "FBA 且非 DRAFT 时必填；FBM 必须为空。",
    }
    ws = wb.create_sheet("导入模板")
    sheet_setup(ws, headers, comments)
    add_enum_validation(ws, "listing_role", ["STANDALONE", "PARENT", "CHILD"])
    add_enum_validation(ws, "fulfillment_channel", ["FBA", "FBM"])
    add_enum_validation(ws, "listing_status", ["DRAFT", "SYNCING", "ACTIVE", "RESTRICTED", "DELETED"])
    good = wb.create_sheet("正确样例")
    sheet_setup(good, headers)
    add_rows(good, [
        ["TEST-STORE-01", "ATVPDKIKX0DER", "TEST-PARENT-01", "B0TESTP001", "PARENT", "TEST-P001", None, None, "ColorSize", None, None, "ACTIVE"],
        ["TEST-STORE-01", "ATVPDKIKX0DER", "TEST-CHILD-01", "B0TESTC001", "CHILD", "TEST-P001", "TEST-SKU001", "TEST-PARENT-01", "ColorSize", "X00TEST001", "FBA", "ACTIVE"],
        ["TEST-STORE-02", "ATVPDKIKX0DER", "TEST-ONLY-02", "B0TESTC001", "STANDALONE", "TEST-P001", "TEST-SKU001", None, None, None, "FBM", "ACTIVE"],
        ["TEST-STORE-03", "ATVPDKIKX0DER", "TEST-ONLY-03", "B0TESTC002", "STANDALONE", "TEST-P001", "TEST-SKU002", None, None, "X00TEST003", "FBA", "ACTIVE"],
        ["TEST-STORE-04", "ATVPDKIKX0DER", "TEST-DRAFT-04", None, "STANDALONE", "TEST-P001", "TEST-SKU002", None, None, None, "FBM", "DRAFT"],
        ["TEST-STORE-05", "ATVPDKIKX0DER", "TEST-ONLY-05", "B0TESTC003", "STANDALONE", "TEST-P001", "TEST-SKU002", None, None, None, "FBM", "RESTRICTED"],
    ])
    bad = wb.create_sheet("错误样例")
    sheet_setup(bad, headers + ["_预期错误"])
    add_rows(bad, [
        ["UNKNOWN-STORE", "ATVPDKIKX0DER", "BAD-01", "SHORT", "STANDALONE", "TEST-P001", "TEST-SKU001", None, None, None, "FBM", "ACTIVE", "店铺引用不存在；ASIN 格式错误"],
        ["TEST-STORE-01", "ATVPDKIKX0DER", "BAD-PARENT", "B0BADP0001", "PARENT", "TEST-P001", "TEST-SKU001", None, "Color", None, "FBA", "ACTIVE", "父体不得关联 SKU/履约"],
        ["TEST-STORE-01", "ATVPDKIKX0DER", "BAD-CHILD", "B0BADC0001", "CHILD", "TEST-P001", None, None, None, None, "FBA", "ACTIVE", "子体缺 SKU/父体/主题/FNSKU"],
        ["TEST-STORE-02", "ATVPDKIKX0DER", "BAD-FBM", "B0BADF0001", "STANDALONE", "TEST-P001", "TEST-SKU001", None, None, "X00BAD0001", "FBM", "ACTIVE", "FBM 不得填写 FNSKU"],
    ])
    add_enum_sheet(wb, {"listing_role": ["STANDALONE", "PARENT", "CHILD"], "fulfillment_channel": ["FBA", "FBM"], "listing_status": ["DRAFT", "SYNCING", "ACTIVE", "RESTRICTED", "DELETED"]})
    save_and_check(wb, "03_store_listings_v1.0.xlsx", ["填写说明", "导入模板", "正确样例", "错误样例", "枚举值"])


def material_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "04 辅料主档导入模板",
        "产品/SKU 可并行；必须先于 05 SKU-辅料关系与 06 期初库存",
        "material_code（大小写不敏感全局唯一）",
        ["material_code 由业务手工提供，系统不得自动生成。", "图片只能填写批准对象存储 URL/引用。"],
    )
    add_machine_metadata(wb, "AUXILIARY_MATERIALS")
    headers = ["material_code", "material_name", "material_type", "specification", "base_unit", "image_url", "status", "remark"]
    ws = wb.create_sheet("导入模板")
    sheet_setup(ws, headers)
    add_enum_validation(ws, "material_type", ["LABEL", "BOX", "BAG", "MANUAL", "INSERT", "OTHER"])
    add_enum_validation(ws, "base_unit", ["PCS", "SET", "ROLL", "SHEET", "M", "KG"])
    add_enum_validation(ws, "status", ["ACTIVE", "INACTIVE"])
    good = wb.create_sheet("正确样例")
    sheet_setup(good, headers)
    add_rows(good, [
        ["TEST-MAT001", "测试标签（非业务数据）", "LABEL", "TEST-SPEC", "PCS", None, "ACTIVE", "仅用于校验"],
        ["TEST-MAT002", "测试包装盒（非业务数据）", "BOX", "TEST-SPEC", "PCS", "https://example.invalid/test-object", "ACTIVE", "仅用于校验"],
    ])
    bad = wb.create_sheet("错误样例")
    sheet_setup(bad, headers + ["_预期错误"])
    add_rows(bad, [
        [None, "缺编码", "LABEL", None, "PCS", None, "ACTIVE", None, "material_code 必填"],
        ["TEST-MAT001", "重复 A", "LABEL", None, "PCS", None, "ACTIVE", None, "与正确样例重复"],
        ["test-mat001", "重复 B", "UNKNOWN", None, "BOXES", "file:///tmp/test.jpg", "UNKNOWN", None, "大小写不敏感重复；枚举/图片/status 无效"],
    ])
    add_enum_sheet(wb, {"material_type": ["LABEL", "BOX", "BAG", "MANUAL", "INSERT", "OTHER"], "base_unit": ["PCS", "SET", "ROLL", "SHEET", "M", "KG"], "status": ["ACTIVE", "INACTIVE"]})
    save_and_check(wb, "04_auxiliary_materials_v1.0.xlsx", ["填写说明", "导入模板", "正确样例", "错误样例", "枚举值"])


def relation_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "05 SKU—辅料关系导入模板",
        "先导入 02 SKU 与 04 辅料",
        "sku_code + material_code",
        ["关系必须挂内部 SKU，不得填写 ASIN/Seller SKU。", "unit 必须与辅料 base_unit 一致；MVP 不做单位换算。"],
    )
    add_machine_metadata(wb, "PRODUCT_AUXILIARY_MATERIALS")
    headers = ["sku_code", "material_code", "standard_quantity", "unit", "loss_rate", "enabled", "remark"]
    ws = wb.create_sheet("导入模板")
    sheet_setup(ws, headers)
    add_enum_validation(ws, "unit", ["PCS", "SET", "ROLL", "SHEET", "M", "KG"])
    add_enum_validation(ws, "enabled", ["TRUE", "FALSE"])
    add_decimal_validation(ws, "standard_quantity", "greaterThan", "0")
    add_decimal_validation(ws, "loss_rate", "between", "0", "1")
    good = wb.create_sheet("正确样例")
    sheet_setup(good, headers)
    add_rows(good, [
        ["TEST-SKU001", "TEST-MAT001", 1, "PCS", 0.02, "TRUE", "测试关系"],
        ["TEST-SKU001", "TEST-MAT002", 1, "PCS", 0, "TRUE", "测试关系"],
        ["TEST-SKU002", "TEST-MAT001", 2, "PCS", 0.03, "TRUE", "应用到全部 SKU 后独立调整示例"],
    ])
    bad = wb.create_sheet("错误样例")
    sheet_setup(bad, headers + ["_预期错误"])
    add_rows(bad, [
        ["UNKNOWN-SKU", "TEST-MAT001", 1, "PCS", 0, "TRUE", None, "SKU 引用不存在"],
        ["TEST-SKU001", "UNKNOWN-MAT", 1, "PCS", 0, "TRUE", None, "辅料引用不存在"],
        ["TEST-SKU001", "TEST-MAT001", 0, "KG", 1.2, "YES", None, "关系重复；用量、单位、损耗率、布尔值无效"],
    ])
    add_enum_sheet(wb, {"unit": ["PCS", "SET", "ROLL", "SHEET", "M", "KG"], "enabled": ["TRUE", "FALSE"]})
    save_and_check(wb, "05_sku_auxiliary_materials_v1.0.xlsx", ["填写说明", "导入模板", "正确样例", "错误样例", "枚举值"])


def inventory_workbook() -> None:
    wb = Workbook()
    add_instructions(
        wb,
        "06 单仓辅料期初库存导入模板",
        "先导入 04 辅料；期初导入经预览确认后生成不可编辑 OPENING 流水",
        "warehouse_code + material_code；opening_reference 在仓库内唯一",
        ["MVP 只允许系统配置的唯一启用仓库。", "余额纠错须新增反向调整，不得编辑/删除流水。"],
    )
    add_machine_metadata(wb, "AUXILIARY_MATERIAL_OPENING_INVENTORY")
    headers = ["warehouse_code", "material_code", "opening_quantity", "safety_stock", "as_of_date", "opening_reference", "remark"]
    ws = wb.create_sheet("导入模板")
    sheet_setup(ws, headers)
    add_decimal_validation(ws, "opening_quantity", "greaterThanOrEqual", "0")
    add_decimal_validation(ws, "safety_stock", "greaterThanOrEqual", "0")
    good = wb.create_sheet("正确样例")
    sheet_setup(good, headers)
    add_rows(good, [
        ["TEST-WH-01", "TEST-MAT001", 100, 20, "2026-09-12", "TEST-OPEN-001", "合成测试数据"],
        ["TEST-WH-01", "TEST-MAT002", 50, 10, "2026-09-12", "TEST-OPEN-002", "合成测试数据"],
    ])
    bad = wb.create_sheet("错误样例")
    sheet_setup(bad, headers + ["_预期错误"])
    add_rows(bad, [
        ["UNKNOWN-WH", "TEST-MAT001", 10, 1, "2026-09-12", "TEST-OPEN-BAD1", None, "仓库不是系统固定单仓"],
        ["TEST-WH-01", "UNKNOWN-MAT", 10, 1, "2026-09-12", "TEST-OPEN-BAD2", None, "辅料引用不存在"],
        ["TEST-WH-01", "TEST-MAT001", -1, -2, "2099-01-01", "TEST-OPEN-001", None, "重复期初；数量/安全库存为负；日期晚于提交日；依据号重复"],
        ["TEST-WH-01", "TEST-MAT002", 10, 1, None, None, None, "期初日期和依据号必填"],
    ])
    save_and_check(wb, "06_initial_inventory_v1.0.xlsx", ["填写说明", "导入模板", "正确样例", "错误样例"])


def main() -> None:
    product_workbook()
    sku_workbook()
    listing_workbook()
    material_workbook()
    relation_workbook()
    inventory_workbook()
    print(f"generated 6 workbooks in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
